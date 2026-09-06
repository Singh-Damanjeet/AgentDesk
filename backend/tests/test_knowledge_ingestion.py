from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401

from app.api.routes import knowledge as knowledge_route
from app.db.base import Base
from app.db.session import get_db
from app.knowledge.chunking import TextChunker
from app.knowledge.config import ChunkSettings, MAX_DOCUMENT_SIZE_BYTES
from app.knowledge.errors import KnowledgeProcessingError
from app.knowledge.normalization import TextNormalizer
from app.knowledge.parsers.base import ParsedSection
from app.knowledge.parsers.factory import DocumentParserFactory
from app.knowledge.parsers.pdf import PDFParser
from app.knowledge.storage import DocumentStorage
from app.knowledge.vector_store import LocalVectorStore, VectorRecord
from app.main import app
from app.models.knowledge_chunk import KnowledgeChunk
from app.models.knowledge_document import KnowledgeDocument
from app.services.knowledge_service import KnowledgeService


class FakeEmbeddingService:
    def __init__(self):
        self.calls: list[list[str]] = []

    async def embed_text(self, text: str) -> list[float]:
        return (await self.embed_batch([text]))[0]

    async def embed_batch(self, texts: list[str]) -> list[list[float]]:
        self.calls.append(list(texts))
        return [
            [float(len(text)), float(index + 1), 0.5]
            for index, text in enumerate(texts)
        ]


class FailingEmbeddingService:
    async def embed_text(self, text: str) -> list[float]:
        del text
        raise RuntimeError("synthetic embedding failure")

    async def embed_batch(self, texts: list[str]) -> list[list[float]]:
        del texts
        raise RuntimeError("synthetic embedding failure")


@pytest.fixture()
def knowledge_client(tmp_path):
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)

    def override_get_db():
        with Session(engine) as session:
            yield session

    embedding_service = FakeEmbeddingService()
    knowledge_service = KnowledgeService(
        storage=DocumentStorage(tmp_path / "knowledge"),
        embedding_service=embedding_service,
    )
    previous_override = app.dependency_overrides.get(get_db)
    previous_service = knowledge_route.knowledge_service
    app.dependency_overrides[get_db] = override_get_db
    knowledge_route.knowledge_service = knowledge_service

    try:
        with TestClient(app) as client:
            yield client, engine, knowledge_service, embedding_service
    finally:
        knowledge_route.knowledge_service = previous_service
        if previous_override is None:
            app.dependency_overrides.pop(get_db, None)
        else:
            app.dependency_overrides[get_db] = previous_override
        Base.metadata.drop_all(engine)
        engine.dispose()


def create_pdf_bytes(text: str) -> bytes:
    escaped_text = (
        text.replace("\\", "\\\\")
        .replace("(", "\\(")
        .replace(")", "\\)")
    )
    stream = (
        "BT /F1 12 Tf 72 720 Td ("
        f"{escaped_text}"
        ") Tj ET"
    ).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        (
            f"<< /Length {len(stream)} >>\nstream\n".encode()
            + stream
            + b"\nendstream"
        ),
    ]

    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]

    for object_index, value in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{object_index} 0 obj\n".encode())
        pdf.extend(value)
        pdf.extend(b"\nendobj\n")

    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())

    pdf.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} "
            f"/Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(pdf)


def create_docx_bytes(tmp_path: Path) -> bytes:
    source_path = tmp_path / "fixture.docx"
    document = Document()
    document.add_heading("Refund policy", level=1)
    document.add_paragraph("Refunds are available within thirty days.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Plan"
    table.cell(0, 1).text = "Refund window"
    table.cell(1, 0).text = "Standard"
    table.cell(1, 1).text = "30 days"
    document.save(source_path)
    return source_path.read_bytes()


def test_parser_factory_supports_pdf_docx_txt_and_markdown(tmp_path):
    pdf_path = tmp_path / "policy.pdf"
    pdf_path.write_bytes(create_pdf_bytes("Refunds are available."))

    docx_path = tmp_path / "policy.docx"
    docx_path.write_bytes(create_docx_bytes(tmp_path))

    txt_path = tmp_path / "policy.txt"
    txt_path.write_text("Refunds are available.\n", encoding="utf-8")

    markdown_path = tmp_path / "policy.md"
    markdown_path.write_text(
        "# Refund policy\n\nRefunds are available.",
        encoding="utf-8",
    )

    pdf_sections = DocumentParserFactory.create("pdf").parse(pdf_path)
    docx_sections = DocumentParserFactory.create("docx").parse(docx_path)
    txt_sections = DocumentParserFactory.create("txt").parse(txt_path)
    markdown_sections = DocumentParserFactory.create("markdown").parse(
        markdown_path
    )

    assert pdf_sections[0].page == 1
    assert "Refunds" in pdf_sections[0].text
    assert any(section.metadata.get("kind") == "table" for section in docx_sections)
    assert txt_sections[0].text.startswith("Refunds")
    assert markdown_sections[0].section == "Refund policy"


def test_pdf_parser_rejects_empty_or_image_only_content(tmp_path):
    empty_pdf = tmp_path / "empty.pdf"
    empty_pdf.write_bytes(create_pdf_bytes(" "))

    with pytest.raises(KnowledgeProcessingError, match="readable text"):
        PDFParser().parse(empty_pdf)

    with pytest.raises(KnowledgeProcessingError):
        PDFParser().parse(tmp_path / "missing.pdf")


def test_normalization_and_chunking_are_deterministic():
    parsed = [
        ParsedSection(
            text="One   two\n\n\nthree four five six seven eight nine.",
            page=2,
            section="Refunds",
        )
    ]
    normalized = TextNormalizer.normalize(parsed)
    chunker = TextChunker(ChunkSettings(target_tokens=5, overlap_tokens=2))

    first = chunker.chunk(normalized)
    second = chunker.chunk(normalized)

    assert first == second
    assert [chunk.token_count for chunk in first] == [5, 5, 3]
    assert all(chunk.page == 2 for chunk in first)
    assert all(chunk.section == "Refunds" for chunk in first)


def test_upload_lists_details_and_stores_batched_vectors(knowledge_client):
    client, engine, knowledge_service, embedding_service = knowledge_client

    response = client.post(
        "/api/knowledge/documents",
        files={
            "file": (
                "refund-policy.txt",
                b"Refunds are available within thirty days.",
                "text/plain",
            )
        },
    )

    assert response.status_code == 201
    payload = response.json()
    assert payload["status"] == "ready"
    assert payload["file_type"] == "txt"
    assert payload["chunk_count"] == 1
    assert len(embedding_service.calls) == 1

    document_id = payload["id"]
    with Session(engine) as db:
        document = db.get(KnowledgeDocument, document_id)
        chunks = db.scalars(
            select(KnowledgeChunk).where(
                KnowledgeChunk.document_id == document_id
            )
        ).all()

        assert document is not None
        assert len(chunks) == 1
        assert chunks[0].chunk_metadata["source_filename"] == (
            "refund-policy.txt"
        )
        assert chunks[0].embedding == [41.0, 1.0, 0.5]

    source_path = knowledge_service.storage.source_path(document_id, "txt")
    assert source_path.is_file()

    listed = client.get("/api/knowledge/documents")
    detail = client.get(f"/api/knowledge/documents/{document_id}")
    assert listed.status_code == 200
    assert listed.json()[0]["id"] == document_id
    assert detail.status_code == 200
    assert detail.json()["chunks"][0]["chunk_index"] == 0
    assert "embedding" not in detail.text


def test_default_local_embeddings_do_not_require_gemini_configuration(
    tmp_path,
    monkeypatch,
):
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    embedding_service = FakeEmbeddingService()
    factory_calls: list[tuple[str | None, str | None]] = []

    def create_embedding_service(
        *,
        provider: str | None = None,
        model: str | None = None,
    ) -> FakeEmbeddingService:
        factory_calls.append((provider, model))
        return embedding_service

    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.setattr(
        "app.services.embedding_configuration_service.EmbeddingFactory.create",
        create_embedding_service,
    )

    service = KnowledgeService(
        storage=DocumentStorage(tmp_path / "knowledge"),
    )

    try:
        with Session(engine) as db:
            document = run_async(
                service.ingest(
                    db,
                    filename="local-only.txt",
                    content_type="text/plain",
                    content=b"Local embeddings do not need Gemini.",
                )
            )

        assert document.status == "ready"
        assert factory_calls == [(None, None)]
        assert len(embedding_service.calls) == 1
    finally:
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_upload_supports_all_document_types(knowledge_client, tmp_path):
    client, _, _, _ = knowledge_client
    uploads = [
        ("policy.pdf", create_pdf_bytes("PDF policy text."), "application/pdf"),
        (
            "policy.docx",
            create_docx_bytes(tmp_path),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        ("policy.txt", b"Plain text policy.", "text/plain"),
        ("policy.md", b"# Policy\n\nMarkdown policy.", "text/markdown"),
    ]

    for filename, content, content_type in uploads:
        response = client.post(
            "/api/knowledge/documents",
            files={"file": (filename, content, content_type)},
        )
        assert response.status_code == 201
        assert response.json()["status"] == "ready"


def test_invalid_uploads_are_rejected_before_persistence(knowledge_client):
    client, engine, _, _ = knowledge_client
    invalid_uploads = [
        ("policy.exe", b"content", "application/octet-stream"),
        ("policy.pdf", b"content", "text/plain"),
        ("", b"content", "text/plain"),
        ("policy.txt", b"", "text/plain"),
        (
            "policy.txt",
            b"x" * (MAX_DOCUMENT_SIZE_BYTES + 1),
            "text/plain",
        ),
        ("../../secret.key", b"content", "application/octet-stream"),
    ]

    for filename, content, content_type in invalid_uploads:
        response = client.post(
            "/api/knowledge/documents",
            files={"file": (filename, content, content_type)},
        )
        assert response.status_code in {400, 422}

    with Session(engine) as db:
        assert db.scalar(select(KnowledgeDocument)) is None


def test_corrupt_documents_are_recorded_as_failed(knowledge_client):
    client, engine, _, _ = knowledge_client

    pdf_response = client.post(
        "/api/knowledge/documents",
        files={"file": ("corrupt.pdf", b"not a PDF", "application/pdf")},
    )
    docx_response = client.post(
        "/api/knowledge/documents",
        files={"file": ("corrupt.docx", b"not a DOCX", "application/octet-stream")},
    )

    assert pdf_response.status_code == 201
    assert docx_response.status_code == 201
    assert pdf_response.json()["status"] == "failed"
    assert docx_response.json()["status"] == "failed"
    assert "read" in pdf_response.json()["error"].lower()
    assert "DOCX" in docx_response.json()["error"]

    with Session(engine) as db:
        assert db.scalar(select(KnowledgeChunk)) is None


def test_reindex_replaces_chunks_without_duplication(knowledge_client):
    client, engine, knowledge_service, _ = knowledge_client
    response = client.post(
        "/api/knowledge/documents",
        files={"file": ("faq.md", b"# FAQ\n\nFirst answer.", "text/markdown")},
    )
    document_id = response.json()["id"]
    source_path = knowledge_service.storage.source_path(document_id, "markdown")
    source_path.write_text(
        "# FAQ\n\nFirst answer.\n\nSecond answer.",
        encoding="utf-8",
    )

    reindex_response = client.post(
        f"/api/knowledge/documents/{document_id}/reindex"
    )

    assert reindex_response.status_code == 200
    assert reindex_response.json()["status"] == "ready"

    with Session(engine) as db:
        chunks = db.scalars(
            select(KnowledgeChunk)
            .where(KnowledgeChunk.document_id == document_id)
            .order_by(KnowledgeChunk.chunk_index)
        ).all()
        assert len(chunks) == reindex_response.json()["chunk_count"]
        assert [chunk.chunk_index for chunk in chunks] == list(range(len(chunks)))


def test_failed_embedding_marks_document_failed_and_cleans_partial_vectors(
    tmp_path,
):
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    service = KnowledgeService(
        storage=DocumentStorage(tmp_path / "knowledge"),
        embedding_service=FailingEmbeddingService(),
    )

    try:
        with Session(engine) as db:
            document = run_async(
                service.ingest(
                    db,
                    filename="failure.txt",
                    content_type="text/plain",
                    content=b"This will fail during embedding.",
                )
            )
            assert document.status == "failed"
            assert document.error == "Document processing failed."
            assert db.scalar(select(KnowledgeChunk)) is None
    finally:
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_delete_removes_document_chunks_vectors_and_source(knowledge_client):
    client, engine, knowledge_service, _ = knowledge_client
    response = client.post(
        "/api/knowledge/documents",
        files={"file": ("delete.txt", b"Delete me.", "text/plain")},
    )
    document_id = response.json()["id"]
    source_path = knowledge_service.storage.source_path(document_id, "txt")
    assert source_path.exists()

    delete_response = client.delete(f"/api/knowledge/documents/{document_id}")

    assert delete_response.status_code == 204
    assert not source_path.parent.exists()
    assert client.get(f"/api/knowledge/documents/{document_id}").status_code == 404

    with Session(engine) as db:
        assert db.get(KnowledgeDocument, document_id) is None
        assert db.scalar(select(KnowledgeChunk)) is None


def test_local_vector_store_add_search_and_delete():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)

    try:
        with Session(engine) as db:
            document_id = "00000000-0000-0000-0000-000000000001"
            db.add(
                KnowledgeDocument(
                    id=document_id,
                    filename="vectors.txt",
                    original_filename="vectors.txt",
                    file_type="txt",
                    file_size=10,
                    status="ready",
                )
            )
            db.commit()
            store = LocalVectorStore(db)
            store.add(
                [
                    VectorRecord(
                        id="00000000-0000-0000-0000-000000000002",
                        document_id=document_id,
                        chunk_index=0,
                        content="Vector content",
                        page_number=None,
                        section=None,
                        token_count=2,
                        char_start=0,
                        char_end=14,
                        embedding=[1.0, 0.0],
                        metadata={"source_filename": "vectors.txt"},
                    )
                ]
            )
            db.commit()

            results = store.search([1.0, 0.0])
            assert len(results) == 1
            assert results[0].score == 1.0

            store.delete_document(document_id)
            db.commit()
            assert store.search([1.0, 0.0]) == []
    finally:
        Base.metadata.drop_all(engine)
        engine.dispose()


def run_async(coroutine):
    import asyncio

    return asyncio.run(coroutine)
