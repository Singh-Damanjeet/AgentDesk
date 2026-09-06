from pathlib import Path
from zipfile import BadZipFile, ZipFile

from app.knowledge.errors import KnowledgeProcessingError
from app.knowledge.parsers.base import ParsedSection


class DocxParser:
    def parse(self, source_path: Path) -> list[ParsedSection]:
        try:
            with ZipFile(source_path) as archive:
                names = archive.namelist()
                if (
                    "[Content_Types].xml" not in names
                    or "word/document.xml" not in names
                ):
                    raise KnowledgeProcessingError(
                        "The uploaded file is not a valid DOCX document."
                    )

                if any(
                    name.lower().endswith("vbaproject.bin")
                    for name in names
                ):
                    raise KnowledgeProcessingError(
                        "Macro-enabled DOCX files are not supported."
                    )

            from docx import Document

            document = Document(str(source_path))
        except ImportError as exc:
            raise KnowledgeProcessingError(
                "DOCX parsing is not available in this installation."
            ) from exc
        except (BadZipFile, OSError) as exc:
            raise KnowledgeProcessingError(
                "Unable to read this DOCX document."
            ) from exc
        except KnowledgeProcessingError:
            raise
        except Exception as exc:
            raise KnowledgeProcessingError(
                "Unable to read this DOCX document."
            ) from exc

        sections: list[ParsedSection] = []
        current_heading: str | None = None

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue

            style_name = getattr(paragraph.style, "name", "") or ""
            is_heading = style_name.lower().startswith("heading")

            if is_heading:
                current_heading = text

            sections.append(
                ParsedSection(
                    text=text,
                    section=current_heading,
                    metadata={
                        "kind": "heading" if is_heading else "paragraph",
                    },
                )
            )

        for table in document.tables:
            rows: list[str] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                cells = [cell for cell in cells if cell]
                if cells:
                    rows.append(" | ".join(cells))

            if rows:
                sections.append(
                    ParsedSection(
                        text="\n".join(rows),
                        section=current_heading,
                        metadata={"kind": "table"},
                    )
                )

        if not sections:
            raise KnowledgeProcessingError(
                "The DOCX document does not contain readable content."
            )

        return sections
